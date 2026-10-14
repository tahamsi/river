#!/usr/bin/env python3
"""Build the pair-safe, human-adjudicated RIVER benchmark release.

The annotation workbook is the decision ledger.  A pair is retained only when
the Pair QA sheet marks it Ready, both row labels are clear, and neither
construction row is rejected.  Human state and active-hypothesis choices
replace the provisional generated labels.  The input corpus is never changed.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


SPLITS = ("train", "dev", "test")
SOURCE_DIRS = ("grounded_nli", "grounded_expired_nli")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return f"sha256:{digest.hexdigest()}"


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    with path.open(encoding="utf-8") as handle:
        return [json.loads(line) for line in handle if line.strip()]


def write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")


def worksheet_rows(workbook: Any, name: str) -> list[dict[str, Any]]:
    sheet = workbook[name]
    values = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(values)]
    return [
        dict(zip(headers, row))
        for row in values
        if any(value is not None and value != "" for value in row)
    ]


def normalise(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().split())


def parse_hypotheses(value: Any) -> dict[str, tuple[str, str, str]]:
    options: dict[str, tuple[str, str, str]] = {}
    for line in str(value or "").splitlines():
        if ":" not in line:
            continue
        label, triple = line.split(":", 1)
        parts = [normalise(part) for part in triple.split("|")]
        if label.strip() in {"H1", "H2"} and len(parts) == 3:
            options[label.strip()] = (parts[0], parts[1], parts[2])
    return options


def selected_hypothesis_id(row: dict[str, Any], source: dict[str, Any]) -> str | None:
    state = normalise(row["Human state"])
    selected = str(row.get("Selected hypothesis") or "").strip()
    if state != "active":
        return None
    if selected not in {"H1", "H2"}:
        raise ValueError(f"Active example {row['Example ID']} has no H1/H2 selection")
    wanted = parse_hypotheses(row["Hypothesis options"]).get(selected)
    matches = [
        item["hypothesis_id"]
        for item in source["hypotheses"]
        if (normalise(item["subject"]), normalise(item["predicate"]), normalise(item["object"])) == wanted
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Could not uniquely map {selected} for {row['Example ID']}: {len(matches)} matches"
        )
    return matches[0]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--data-root", type=Path, default=Path("data"))
    parser.add_argument("--output", type=Path, default=Path("data/grounded_human_frozen"))
    args = parser.parse_args()

    workbook_path = args.workbook.resolve()
    data_root = args.data_root.resolve()
    output = args.output.resolve()
    if output.exists() and any(output.iterdir()):
        raise SystemExit(f"Refusing to overwrite non-empty frozen release: {output}")
    output.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    labels = worksheet_rows(workbook, "Label Review")
    constructions = worksheet_rows(workbook, "Construction Audit")
    pair_qa = worksheet_rows(workbook, "Pair QA")

    if len(labels) != 1780 or len(constructions) != 1780 or len(pair_qa) != 890:
        raise ValueError("Workbook dimensions differ from the adjudicated 1,780-row release")

    label_by_id = {str(row["Example ID"]): row for row in labels}
    construction_by_id = {str(row["Example ID"]): row for row in constructions}
    if len(label_by_id) != len(labels) or len(construction_by_id) != len(constructions):
        raise ValueError("Duplicate example IDs in the annotation workbook")

    ready_pairs = {
        str(row["Pair ID"])
        for row in pair_qa
        if str(row.get("Pair status") or "").strip() == "Ready"
    }
    excluded_pair_rows = [row for row in pair_qa if str(row["Pair ID"]) not in ready_pairs]
    if len(ready_pairs) != 751 or len(excluded_pair_rows) != 139:
        raise ValueError(
            f"Expected 751 Ready and 139 excluded pairs; got {len(ready_pairs)} and {len(excluded_pair_rows)}"
        )

    corpus: dict[str, dict[str, Any]] = {}
    input_hashes: dict[str, str] = {}
    for source_dir in SOURCE_DIRS:
        for split in SPLITS:
            path = data_root / source_dir / f"river_bench_{split}.jsonl"
            input_hashes[str(path.relative_to(data_root.parent))] = sha256(path)
            for row in read_jsonl(path):
                corpus.setdefault(str(row["example_id"]), row)

    missing = sorted(set(label_by_id) - set(corpus))
    if missing:
        raise ValueError(f"{len(missing)} workbook examples are absent from source corpora: {missing[:5]}")

    frozen: list[dict[str, Any]] = []
    ledger: list[dict[str, Any]] = []
    for example_id, annotation in label_by_id.items():
        pair_id = str(annotation["Pair ID"])
        construction = construction_by_id[example_id]
        included = pair_id in ready_pairs
        if included:
            if annotation["Human label status"] != "Clear label":
                raise ValueError(f"Ready pair contains unresolved label: {example_id}")
            if construction["Disposition"] == "Reject":
                raise ValueError(f"Ready pair contains rejected construction: {example_id}")
            source = corpus[example_id]
            revised = copy.deepcopy(source)
            revised["gold_state"] = normalise(annotation["Human state"])
            revised["gold_hypothesis_id"] = selected_hypothesis_id(annotation, source)
            metadata = revised.setdefault("transformation_metadata", {})
            metadata.update(
                {
                    "annotation_status": "human_adjudicated_frozen",
                    "human_review_status": "complete",
                    "human_label_confidence": annotation["Confidence (1-5)"],
                    "human_label_record_checksum": annotation["Record checksum"],
                    "construction_record_checksum": construction["Record checksum"],
                    "construction_disposition": construction["Disposition"],
                    "freeze_policy": "pair_qa_ready_v1",
                }
            )
            frozen.append(revised)
        ledger.append(
            {
                "example_id": example_id,
                "pair_id": pair_id,
                "split": annotation["Split"],
                "included": included,
                "human_state": normalise(annotation["Human state"]),
                "human_label_status": annotation["Human label status"],
                "construction_disposition": construction["Disposition"],
                "pair_status": next(
                    str(row["Pair status"]) for row in pair_qa if str(row["Pair ID"]) == pair_id
                ),
                "label_record_checksum": annotation["Record checksum"],
                "construction_record_checksum": construction["Record checksum"],
            }
        )

    frozen.sort(key=lambda row: (SPLITS.index(row["split"]), row["pair_id"], row["example_id"]))
    if len(frozen) != 1502:
        raise ValueError(f"Expected 1,502 frozen examples; got {len(frozen)}")
    pair_sizes = Counter(row["pair_id"] for row in frozen)
    if set(pair_sizes.values()) != {2}:
        raise ValueError("Frozen corpus contains incomplete pairs")

    output_files: dict[str, str] = {}
    for split in SPLITS:
        path = output / f"river_bench_{split}.jsonl"
        write_jsonl(path, (row for row in frozen if row["split"] == split))
        output_files[path.name] = sha256(path)
    combined_path = output / "river_bench_1502.jsonl"
    write_jsonl(combined_path, frozen)
    output_files[combined_path.name] = sha256(combined_path)
    ledger_path = output / "adjudication_ledger.jsonl"
    write_jsonl(ledger_path, ledger)
    output_files[ledger_path.name] = sha256(ledger_path)

    excluded_by_reason = Counter()
    for row in excluded_pair_rows:
        dispositions = {str(row.get("Construction disposition 1")), str(row.get("Construction disposition 2"))}
        statuses = {str(row.get("Label status 1")), str(row.get("Label status 2"))}
        if "Reject" in dispositions:
            excluded_by_reason["construction_reject"] += 1
        elif "Needs adjudication" in statuses:
            excluded_by_reason["unresolved_label"] += 1
        else:
            excluded_by_reason["pair_label_inconsistency"] += 1

    manifest = {
        "schema_version": "river-bench-grounded-v4-human-frozen",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "freeze_policy": {
            "unit": "pair",
            "include_when": "Pair QA status is Ready, both labels are clear, and construction is not rejected",
            "labels": "Human state and active hypothesis replace provisional generated gold",
            "source_scores": "Existing measured NLI verifier scores are retained; no new source data are introduced",
        },
        "annotation_workbook": {
            "filename": workbook_path.name,
            "sha256": sha256(workbook_path),
            "label_rows": len(labels),
            "pair_rows": len(pair_qa),
        },
        "counts": {
            "input_examples": len(labels),
            "input_pairs": len(pair_qa),
            "frozen_examples": len(frozen),
            "frozen_pairs": len(pair_sizes),
            "excluded_examples": len(labels) - len(frozen),
            "excluded_pairs": len(excluded_pair_rows),
            "unresolved_input_rows": sum(row["Human label status"] != "Clear label" for row in labels),
            "by_split": dict(Counter(row["split"] for row in frozen)),
            "by_state": dict(Counter(row["gold_state"] for row in frozen)),
            "by_intervention": dict(Counter(row["intervention"] for row in frozen)),
            "excluded_pairs_by_primary_reason": dict(excluded_by_reason),
        },
        "input_sha256": input_hashes,
        "output_sha256": output_files,
    }
    manifest_path = output / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(manifest["counts"], indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
